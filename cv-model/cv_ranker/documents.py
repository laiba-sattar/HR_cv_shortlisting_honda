"""
Text extraction from CV/JD files.

Unglamorous, and the single most common source of real-world failure: if the
text comes out wrong, everything downstream is garbage regardless of how good
the model is. Fail loudly and per-file rather than silently returning junk.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".doc", ".txt", ".md"}

# Below this many characters we assume extraction failed (e.g. a scanned PDF
# with no text layer) and OCR is needed.
MIN_USABLE_CHARS = 120


class DocumentError(RuntimeError):
    pass


@dataclass
class ExtractedDocument:
    path: str
    text: str
    method: str           # "pdf-text" | "docx" | "plain" | "ocr"
    char_count: int
    warning: str | None = None


def _read_pdf(path: Path) -> tuple[str, str]:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        try:
            import pdfplumber
        except ImportError as e:
            raise DocumentError(
                "Install a PDF reader: pip install pymupdf (or pdfplumber)"
            ) from e
        with pdfplumber.open(str(path)) as pdf:
            return "\n".join(p.extract_text() or "" for p in pdf.pages), "pdf-text"
    doc = fitz.open(str(path))
    try:
        return "\n".join(page.get_text() for page in doc), "pdf-text"
    finally:
        doc.close()


def _ocr_pdf(path: Path) -> str:
    """OCR fallback for scanned CVs. Requires tesseract to be installed."""
    try:
        import fitz
        import pytesseract
        from PIL import Image
    except ImportError as e:
        raise DocumentError(
            "Scanned PDF needs OCR: pip install pymupdf pytesseract pillow, "
            "and install the tesseract binary."
        ) from e
    import io

    doc = fitz.open(str(path))
    out = []
    try:
        for page in doc:
            pix = page.get_pixmap(dpi=300)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            out.append(pytesseract.image_to_string(img))
    finally:
        doc.close()
    return "\n".join(out)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as e:
        raise DocumentError("pip install python-docx") from e
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    # Many CVs lay everything out inside tables — skipping these loses most
    # of the content, which is an easy failure to miss.
    for table in d.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def read_document(path: str | Path, allow_ocr: bool = True) -> ExtractedDocument:
    p = Path(path)
    if not p.exists():
        raise DocumentError(f"File not found: {p}")
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED:
        raise DocumentError(f"Unsupported file type {suffix!r} for {p.name}")

    warning = None

    if suffix == ".pdf":
        text, method = _read_pdf(p)
        if len(text.strip()) < MIN_USABLE_CHARS and allow_ocr:
            log.info("%s looks like a scan — running OCR", p.name)
            text = _ocr_pdf(p)
            method = "ocr"
            warning = "Text recovered by OCR; extraction may be less reliable."
    elif suffix in (".docx", ".doc"):
        if suffix == ".doc":
            raise DocumentError(
                f"{p.name}: legacy .doc is not supported. Convert to .docx first "
                "(e.g. with LibreOffice: soffice --headless --convert-to docx)."
            )
        text, method = _read_docx(p), "docx"
    else:
        text, method = p.read_text(encoding="utf-8", errors="replace"), "plain"

    text = "\n".join(line.rstrip() for line in text.splitlines())
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    text = text.strip()

    if len(text) < MIN_USABLE_CHARS:
        raise DocumentError(
            f"{p.name}: only {len(text)} characters of text recovered — "
            "the file is likely empty, corrupted, or an image-only scan without OCR."
        )

    return ExtractedDocument(str(p), text, method, len(text), warning)
